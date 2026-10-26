from pathlib import Path

from docx import Document

from konwerter.errors import ConversionError


def txt_to_docx(source: Path, target: Path) -> None:
    try:
        with open(source, encoding="utf-8-sig") as file:
            text = file.read()
    except UnicodeDecodeError:
        raise ConversionError("Plik nie jest zapisany w kodowaniu UTF-8. Zapisz go jako UTF-8 i spróbuj ponownie.")

    if text.strip() == "":
        raise ConversionError("Plik TXT jest pusty. Wybierz plik z tekstem.")

    # empty line between texts means a new paragraph
    paragraphs = text.split("\n\n")

    document = Document()
    for paragraph in paragraphs:
        # remove spaces and new lines at the start and end
        paragraph = paragraph.strip()
        if paragraph == "":
            continue
        # new line inside a paragraph stays as a line break
        document.add_paragraph(paragraph)
    document.save(target)


def docx_to_txt(source: Path, target: Path) -> None:
    try:
        document = Document(source)
    except Exception:
        # any problem while opening means the file is broken
        raise ConversionError("Nie udało się otworzyć pliku DOCX. Plik jest uszkodzony albo nie jest dokumentem Word.")

    paragraphs = []
    for paragraph in document.paragraphs:
        # skip empty paragraphs
        if paragraph.text.strip() == "":
            continue
        paragraphs.append(paragraph.text)

    # empty line between paragraphs, the same as in txt_to_docx
    text = "\n\n".join(paragraphs)

    with open(target, "w", encoding="utf-8") as file:
        file.write(text + "\n")